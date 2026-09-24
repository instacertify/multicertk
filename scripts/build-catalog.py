#!/usr/bin/env python3
"""Build the interlinked Certko catalogue JSON from the uploaded Excel library."""

from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path("/home/ubuntu/.cursor/projects/workspace/uploads")
OUT = Path("/workspace/src/data/generated")
OUT.mkdir(parents=True, exist_ok=True)

V5 = ROOT / "BIS_Master_HSN_Category_QCO_Status_v5_SchemeI_and_II_Updated_34ae.xlsx"
LABS = ROOT / "BIS_Labs_Category_Scope_Wise_130d.xlsx"
BEE = ROOT / "BEE_Star_Label_Master_Single_Sheet_2026_b8f5.xlsx"
GMARK = ROOT / "GMark_Product_Categories_Standards_IECEE_GSO_520f.xlsx"
EU_TEST = ROOT / "EU_Sector_Testing_and_Notified_Body_Matrix_482b.xlsx"
EU_MANDATE = ROOT / "EU_Sector_Wise_Mandates_DG_GROW_1075.xlsx"
LIMS = ROOT / "BIS_LIMS_200K_Full_Scale_Directory_5b22.xlsx"


def slugify(value: object, fallback: str = "item") -> str:
    text = str(value or "").lower()
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return (text[:96] or fallback)


def unique_slug(base: str, used: set[str]) -> str:
    slug = base or "item"
    i = 2
    while slug in used:
        slug = f"{base[:90]}-{i}"
        i += 1
    used.add(slug)
    return slug


def norm_standard(value: object) -> str:
    text = str(value or "").upper()
    text = text.replace("：", ":")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def standard_key(value: object) -> str:
    text = norm_standard(value)
    text = text.replace("INDIAN STANDARD", "IS")
    m = re.search(r"(IS|IEC|IS/IEC|EN|GSO|ER)\s*([0-9]+[A-Z0-9./-]*)", text)
    if not m:
        return re.sub(r"[^A-Z0-9]+", "", text)[:24]
    prefix, num = m.group(1), m.group(2)
    part = ""
    pm = re.search(r"PART\s*([0-9A-Z]+)", text)
    if pm:
        part = f"P{pm.group(1)}"
    sec = re.search(r"SEC(?:TION)?\s*([0-9A-Z]+)", text)
    if sec:
        part += f"S{sec.group(1)}"
    return f"{prefix}{num}{part}".replace(" ", "")


def qco_bucket(status: object) -> str:
    text = str(status or "").lower()
    if "upcoming" in text or ("notified" in text and "due" in text) or "deferred" in text:
        return "upcoming"
    if "mandatory" in text or "qco in force" in text or "crs" in text:
        return "mandatory"
    return "voluntary"


def parse_inr(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).replace(",", "").replace("Rs", "").replace("₹", "")
    nums = re.findall(r"\d+", text)
    if not nums:
        return None
    return int(nums[0])


def parse_range(value: object) -> tuple[int, int]:
    if value is None:
        return 5000, 650000
    text = str(value).replace(",", "").replace("*", "")
    nums = [int(n) for n in re.findall(r"\d+", text)]
    if len(nums) >= 2:
        return nums[0], nums[1]
    if len(nums) == 1:
        return nums[0], nums[0]
    return 5000, 650000


def category_slug(name: object) -> str:
    return slugify(name, "others")


def load_rows(path: Path, sheet: str, header_row: int = 1):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    headers = None
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == header_row:
            headers = [str(c).strip() if c is not None else f"col{idx}" for idx, c in enumerate(row)]
            continue
        if headers is None or not any(row):
            continue
        yield dict(zip(headers, row))
    wb.close()


def build_products():
    used: set[str] = set()
    products = []
    by_key: dict[str, list[str]] = defaultdict(list)

    for row in load_rows(V5, "ISI Marking Fees (Scheme-I)", 1):
        is_no = row.get("IS No")
        title = row.get("IS Title (Product)")
        if not is_no or not title or str(is_no).startswith("S.No"):
            continue
        standard = str(is_no).strip()
        slug = unique_slug(slugify(f"{title} {standard}"), used)
        key = standard_key(standard)
        lo, hi = parse_range(row.get("Testing Charges (Rs, Indicative)*"))
        product = {
            "slug": slug,
            "name": str(title).strip(),
            "standard": standard,
            "standardKey": key,
            "hsn": str(row.get("HSN Code (8-digit, ITC-HS)#") or row.get("HSN Code (4-digit)") or "").replace(".0", ""),
            "hsn4": str(row.get("HSN Code (4-digit)") or "").replace(".0", ""),
            "category": str(row.get("Broad Category") or "Others").strip(),
            "categorySlug": category_slug(row.get("Broad Category") or "Others"),
            "schemeSlugs": ["bis"],
            "schemeLabel": "ISI Mark Licence",
            "qcoStatus": qco_bucket(row.get("QCO Status")),
            "qcoLabel": str(row.get("QCO Status") or ""),
            "qcoOrder": str(row.get("QCO / Order & Mandatory From (Notification-Effective)^") or ""),
            "testCostMin": lo,
            "testCostMax": hi,
            "timeline": "8–16 weeks",
            "markingFee": {
                "large": parse_inr(row.get("Marking Fee - Large (Rs)")) or 0,
                "medium": parse_inr(row.get("Marking Fee - Medium (Rs)")) or 0,
                "small": parse_inr(row.get("Marking Fee - Small (Rs) (0.5 x Large)")) or 0,
                "micro": parse_inr(row.get("Marking Fee - Micro (Rs) (0.2 x Large)")) or 0,
            },
            "unit": str(row.get("Unit of Product") or ""),
            "excerpt": f"{title} is mapped to {standard} under BIS Scheme I (ISI). HSN {row.get('HSN Code (8-digit, ITC-HS)#') or 'n/a'}. {row.get('QCO Status') or ''}.",
        }
        products.append(product)
        by_key[key].append(slug)

    for row in load_rows(V5, "CRS Marking Fees (Scheme-II)", 1):
        name = row.get("Product Category")
        standard = row.get("IS Standard")
        if not name or not standard or str(name) == "Product Category":
            continue
        key = standard_key(standard)
        slug = unique_slug(slugify(f"{name} {standard}"), used)
        lo, hi = parse_range(row.get("Testing Charges (Rs, Indicative)*"))
        product = {
            "slug": slug,
            "name": str(name).strip(),
            "standard": str(standard).strip(),
            "standardKey": key,
            "hsn": str(row.get("HSN Code (8-digit, ITC-HS)#") or "").replace(".0", ""),
            "hsn4": str(row.get("HSN Code (4-digit)") or "").replace(".0", ""),
            "category": str(row.get("Broad Category") or "Electrical & Electronics").strip(),
            "categorySlug": category_slug(row.get("Broad Category") or "electrical-and-electronics"),
            "schemeSlugs": ["bis"],
            "schemeLabel": "CRS Registration",
            "qcoStatus": qco_bucket(row.get("QCO Status")),
            "qcoLabel": str(row.get("QCO Status") or ""),
            "qcoOrder": str(row.get("Order & Mandatory From^") or ""),
            "testCostMin": lo,
            "testCostMax": hi,
            "timeline": "6–10 weeks",
            "markingFee": {
                "large": parse_inr(row.get("Marking Fee - Large (Rs)")) or 0,
                "medium": 0,
                "small": parse_inr(row.get("Marking Fee - Small (Rs)")) or 0,
                "micro": parse_inr(row.get("Marking Fee - Micro (Rs)")) or 0,
            },
            "unit": "",
            "excerpt": f"{name} requires BIS CRS against {standard}. {row.get('Remarks') or ''}",
        }
        products.append(product)
        by_key[key].append(slug)

    return products, by_key


def build_labs_and_scopes():
    used: set[str] = set()
    labs: dict[str, dict] = {}
    scopes: list[dict] = []
    by_key: dict[str, list[dict]] = defaultdict(list)

    for row in load_rows(LABS, "Category-Scope-wise Data", 1):
        name = row.get("Lab Name")
        standard = row.get("IS Standard")
        if not name or not standard:
            continue
        code = str(row.get("Lab Code") or "").replace(".0", "")
        existing = next((s for s, l in labs.items() if code and l.get("code") == code), None)
        if existing:
            slug = existing
        else:
            slug = unique_slug(slugify(f"{name} {row.get('City') or ''} {code}"), used)
            labs[slug] = {
                "slug": slug,
                "name": str(name).strip(),
                "city": str(row.get("City") or "").strip(),
                "state": str(row.get("State") or "").strip(),
                "code": code,
                "phone": str(row.get("Phone") or ""),
                "email": str(row.get("Email") or ""),
                "address": str(row.get("Full Address") or ""),
                "contact": str(row.get("Contact Person") or ""),
                "categorySlugs": [],
                "standardKeys": [],
            }
        cat = category_slug(row.get("Broad Category") or "")
        lab = labs[slug]
        if cat and cat not in lab["categorySlugs"]:
            lab["categorySlugs"].append(cat)
        key = standard_key(standard)
        if key and key not in lab["standardKeys"]:
            lab["standardKeys"].append(key)
        price = parse_inr(row.get("Test Price (INR excl tax)")) or 0
        scope = {
            "labSlug": slug,
            "standard": str(standard).strip(),
            "standardKey": key,
            "productScope": str(row.get("Product / Scope") or ""),
            "category": str(row.get("Broad Category") or ""),
            "categorySlug": cat,
            "price": price,
        }
        scopes.append(scope)
        by_key[key].append(scope)

    return list(labs.values()), scopes, by_key


def build_bee():
    wb = load_workbook(BEE, read_only=True, data_only=True)
    ws = wb.active
    items = []
    used: set[str] = set()
    section = "mandatory"
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        first = str(row[0] or "")
        title = str(row[1] or "")
        if "VOLUNTARY" in first.upper() or "VOLUNTARY" in title.upper() and row[0] is None:
            if "VOLUNTARY" in (str(row[0] or "") + str(row[1] or "")).upper() and not str(row[0]).isdigit():
                section = "voluntary"
                continue
        if first.upper().startswith("B.") or "VOLUNTARY PRODUCTS" in (str(row[0] or "")).upper():
            section = "voluntary"
            continue
        if not str(row[0]).isdigit() if row[0] is not None else True:
            continue
        name = str(row[1] or "").strip()
        if not name:
            continue
        slug = unique_slug(slugify(name), used)
        items.append({
            "slug": slug,
            "name": name,
            "starMandatory": section == "mandatory",
            "regime": str(row[2] or ""),
            "standard": str(row[3] or ""),
            "starTable": str(row[4] or ""),
            "price": str(row[5] or ""),
            "labs": str(row[6] or ""),
            "summary": f"{name} — BEE star labelling ({section}). Tested as per {row[3]}. {row[2]}.",
        })
    wb.close()
    return items


def build_gmark():
    items = []
    used: set[str] = set()
    for row in load_rows(GMARK, "G-Mark Standards", 1):
        name = row.get("Product Category")
        if not name or name == "Product Category":
            continue
        slug = unique_slug(slugify(name), used)
        items.append({
            "slug": slug,
            "name": str(name).strip(),
            "family": str(row.get("Product Family") or ""),
            "standard": str(row.get("Main Standard(s)") or ""),
            "tests": str(row.get("Typical Safety / Performance Tests") or ""),
            "emc": str(row.get("EMC Required") or ""),
            "cb": str(row.get("IECEE CB Report Accepted") or ""),
            "nb": str(row.get("GSO Notified Body Required") or ""),
            "remarks": str(row.get("Remarks") or ""),
            "summary": f"{name} under G-Mark / GSO. Main standard {row.get('Main Standard(s)')}. EMC: {row.get('EMC Required')}.",
        })
    return items


def build_eu():
    mandates = {r.get("Sector (DG GROW Category)"): r for r in load_rows(EU_MANDATE, "EU Sector Mandates", 4)}
    items = []
    used: set[str] = set()
    for row in load_rows(EU_TEST, "Sector Testing & NB Matrix", 4):
        sector = row.get("Sector")
        if not sector or sector == "Sector":
            continue
        slug = unique_slug(slugify(sector), used)
        extra = mandates.get(sector, {})
        items.append({
            "slug": slug,
            "name": str(sector).strip(),
            "mandate": str(row.get("Governing Mandate / Legislation") or extra.get("Governing EU Mandate / Legislation") or ""),
            "legal": str(row.get("Legal Reference") or extra.get("Legal Reference No.") or ""),
            "testing": str(row.get("Applicable Testing Required (Key Standards / Test Areas)") or ""),
            "nb": str(row.get("Notified Body Assessment Mandatory?") or extra.get("Notified Body Required?") or ""),
            "whenNb": str(row.get("When NB Is / Is Not Required") or ""),
            "route": str(row.get("Conformity Route & Marking") or extra.get("CE Marking / Conformity Route") or ""),
            "url": str(extra.get("Official Sector Page (URL)") or ""),
            "updates": str(extra.get("Key Notifications / Recent Updates") or ""),
            "summary": f"{sector}: {row.get('Governing Mandate / Legislation')}. NB: {row.get('Notified Body Assessment Mandatory?')}.",
        })
    return items


def find_related_slug(name: str, standard: str, products: list[dict]) -> str | None:
    key = standard_key(standard)
    if key:
        for product in products:
            if product.get("standardKey") == key:
                return product["slug"]
    tokens = {t for t in slugify(name).split("-") if len(t) > 2}
    best = None
    best_score = 1
    for product in products:
        score = len(tokens & set(product["slug"].split("-")))
        if score > best_score:
            best = product["slug"]
            best_score = score
    return best


def build_qcos(products: list[dict]):
    items = []
    used: set[str] = set()
    by_key: dict[str, list[str]] = defaultdict(list)
    for product in products:
        by_key[product["standardKey"]].append(product["slug"])

    for row in load_rows(V5, "Upcoming QCOs (Jul-2026)", 4):
        product = row.get("Product")
        standard = row.get("Indian Standard")
        if not product:
            continue
        slug = unique_slug(slugify(f"{product} {standard}"), used)
        key = standard_key(standard)
        product_slugs = list(dict.fromkeys(by_key.get(key, [])))
        related = find_related_slug(str(product), str(standard or ""), products)
        if related and related not in product_slugs:
            product_slugs.append(related)
        items.append({
            "slug": slug,
            "name": str(product).strip(),
            "ministry": str(row.get("Ministry / Department") or ""),
            "standard": str(standard or ""),
            "deadline": str(row.get("Enforcement Date") or ""),
            "scheme": str(row.get("Scheme") or ""),
            "status": "upcoming",
            "productSlugs": product_slugs,
            "summary": f"Upcoming QCO: {product} ({standard}) from {row.get('Enforcement Date')} via {row.get('Scheme')}.",
        })
        for product_slug in product_slugs:
            for mapped in products:
                if mapped["slug"] == product_slug:
                    mapped["qcoSlug"] = slug
                    mapped["qcoStatus"] = "upcoming"

    groups: dict[str, list[dict]] = defaultdict(list)
    for product in products:
        order = str(product.get("qcoOrder") or "").strip()
        if not order or "not under" in order.lower():
            continue
        if product.get("qcoSlug"):
            continue
        groups[order].append(product)

    for order, rows in groups.items():
        if len(rows) < 2 and rows[0]["qcoStatus"] == "voluntary":
            continue
        slug = unique_slug(slugify(order)[:80], used)
        status = rows[0]["qcoStatus"]
        items.append({
            "slug": slug,
            "name": order[:160],
            "ministry": "DPIIT / concerned ministry",
            "standard": rows[0]["standard"],
            "deadline": "",
            "scheme": rows[0].get("schemeLabel") or "BIS",
            "status": status,
            "productSlugs": [row["slug"] for row in rows[:80]],
            "summary": f"{order} covers {len(rows)} mapped Indian Standard{'s' if len(rows) != 1 else ''}.",
        })
        for row in rows:
            row["qcoSlug"] = slug

    return items


def enrich_lims(labs: list[dict], scopes_by_key: dict[str, list[dict]], products_by_key: dict[str, list[str]]):
    """Pull extra unique labs from LIMS when a standard has few recognised labs."""
    existing_codes = {l.get("code") for l in labs}
    used = {l["slug"] for l in labs}
    extra_labs = {}
    added = 0
    wb = load_workbook(LIMS, read_only=True, data_only=True)
    ws = wb["Master Scope Directory"]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1 or not row or not row[1]:
            continue
        code = str(row[1])
        standard = row[7]
        key = standard_key(standard)
        if key not in products_by_key:
            continue
        if len(scopes_by_key.get(key, [])) >= 12:
            continue
        name = str(row[2] or "").strip()
        city = str(row[5] or "").strip()
        state = str(row[6] or "").strip()
        price = parse_inr(row[9]) or 0
        if code not in existing_codes and code not in extra_labs:
            slug = unique_slug(slugify(f"{name} {city} {code}"), used)
            extra_labs[code] = {
                "slug": slug,
                "name": name,
                "city": city,
                "state": state,
                "code": code,
                "phone": "",
                "email": "",
                "address": str(row[4] or ""),
                "contact": str(row[3] or ""),
                "categorySlugs": [],
                "standardKeys": [key],
                "source": "lims",
            }
        lab = extra_labs.get(code) or next((l for l in labs if l.get("code") == code), None)
        if not lab:
            continue
        scope = {
            "labSlug": lab["slug"],
            "standard": str(standard),
            "standardKey": key,
            "productScope": str(row[8] or ""),
            "category": "",
            "categorySlug": "",
            "price": price,
        }
        scopes_by_key[key].append(scope)
        added += 1
        if added >= 8000:
            break
    wb.close()
    labs.extend(extra_labs.values())
    return labs, scopes_by_key


def link_products(products, labs, scopes_by_key):
    lab_lookup = {l["slug"]: l for l in labs}
    for product in products:
        key = product["standardKey"]
        matched = scopes_by_key.get(key, [])
        # also try loose number match
        if not matched:
            num = re.sub(r"[^0-9]", "", product["standard"])[:6]
            for k, rows in scopes_by_key.items():
                if num and num in k:
                    matched = rows
                    break
        prices = [s["price"] for s in matched if s.get("price")]
        product["labSlugs"] = list(dict.fromkeys(s["labSlug"] for s in matched))[:24]
        product["labCount"] = len(product["labSlugs"])
        if prices:
            product["testCostMin"] = min(prices)
            product["testCostMax"] = max(prices)
        if "bee" not in product["schemeSlugs"] and re.search(r"IS\s*(1391|374|616|1476|15750)", product["standard"], re.I):
            product["schemeSlugs"].append("bee")
        if product["categorySlug"] in {"electrical-and-electronics", "electrical-appliances", "electronics-it-goods"}:
            for extra in ("ce", "fcc"):
                if extra not in product["schemeSlugs"]:
                    product["schemeSlugs"].append(extra)
        product["countrySlugs"] = ["india"]
        if "ce" in product["schemeSlugs"]:
            product["countrySlugs"].append("european-union")
        if "fcc" in product["schemeSlugs"]:
            product["countrySlugs"].append("united-states")
        product["testSlugs"] = []
    return products


def build_categories(products):
    counts: dict[str, dict] = {}
    for p in products:
        slug = p["categorySlug"]
        counts.setdefault(slug, {"slug": slug, "name": p["category"], "count": 0})
        counts[slug]["count"] += 1
        counts[slug]["name"] = p["category"]
    return [
        {
            "slug": c["slug"],
            "name": c["name"],
            "summary": f"{c['count']} mapped Indian Standards and product records in {c['name']}.",
        }
        for c in sorted(counts.values(), key=lambda x: -x["count"])
    ]


def main():
    print("Building products from v5 master...")
    products, products_by_key = build_products()
    print(f"  {len(products)} products, {len(products_by_key)} standard keys")

    print("Building labs/scopes...")
    labs, scopes, scopes_by_key = build_labs_and_scopes()
    print(f"  {len(labs)} labs, {len(scopes)} scope rows")

    print("Enriching from LIMS where a standard is thin...")
    labs, scopes_by_key = enrich_lims(labs, scopes_by_key, products_by_key)
    print(f"  {len(labs)} labs after LIMS enrich")

    products = link_products(products, labs, scopes_by_key)
    categories = build_categories(products)
    bee = build_bee()
    gmark = build_gmark()
    eu = build_eu()
    for item in bee:
        item["relatedProductSlug"] = find_related_slug(item["name"], item.get("standard") or "", products)
    for item in gmark:
        item["relatedProductSlug"] = find_related_slug(item["name"], item.get("standard") or "", products)
    qcos = build_qcos(products)

    # compact scopes for runtime: only those attached to products
    needed = {p["standardKey"] for p in products}
    compact_scopes = []
    for key, rows in scopes_by_key.items():
        if key not in needed:
            continue
        # keep cheapest + diversity, max 20
        rows_sorted = sorted(rows, key=lambda r: r.get("price") or 10**9)
        seen = set()
        for row in rows_sorted:
            if row["labSlug"] in seen:
                continue
            seen.add(row["labSlug"])
            compact_scopes.append(row)
            if len(seen) >= 20:
                break

    used_labs = {s["labSlug"] for s in compact_scopes}
    for p in products:
        used_labs.update(p.get("labSlugs") or [])
    labs = [l for l in labs if l["slug"] in used_labs]

    payload = {
        "stats": {
            "products": len(products),
            "labs": len(labs),
            "scopes": len(compact_scopes),
            "categories": len(categories),
            "bee": len(bee),
            "gmark": len(gmark),
            "eu": len(eu),
            "qcos": len(qcos),
        },
        "products": products,
        "labs": labs,
        "scopes": compact_scopes,
        "categories": categories,
        "bee": bee,
        "gmark": gmark,
        "eu": eu,
        "qcos": qcos,
    }
    out = OUT / "library.json"
    out.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")))
    print("Wrote", out, "bytes", out.stat().st_size)
    print(json.dumps(payload["stats"], indent=2))


if __name__ == "__main__":
    main()
